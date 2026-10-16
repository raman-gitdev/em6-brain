"""excel-tool: lets the brain look inside an uploaded spreadsheet.

Read-only. Works on files uploaded through the chat (FILES_DIR/<file_id>/).
Output is short plain text, sized for a small model's memory.
"""
import json
import os
import re

from fastapi import FastAPI
from openpyxl.utils import get_column_letter, range_boundaries
from pydantic import BaseModel

from .workbook import Sheet, fmt, load

FILES_DIR = os.getenv("FILES_DIR", "/data/uploads")
PROFILE_MAX_ROWS = 400        # rows read per sheet for the profile
MAX_SHEETS = 12
MAX_COLS = 20
SAMPLE_ROWS = 5
MAX_RANGE_CELLS = 400
MAX_SEARCH_HITS = 30

app = FastAPI(title="excel-tool")

FILE_ID = {"type": "string", "description": "file_id of the attached file"}
TOOLS = [
    {
        "name": "excel_profile",
        "description": "Overview of an attached spreadsheet: sheets, tables, headers, row labels, notes. Call first.",
        "parameters": {"type": "object", "properties": {"file_id": FILE_ID},
                       "required": ["file_id"]},
    },
    {
        "name": "excel_read_range",
        "description": "Read exact cells from one sheet, e.g. A1:H20.",
        "parameters": {"type": "object", "properties": {
            "file_id": FILE_ID,
            "sheet": {"type": "string", "description": "sheet name"},
            "range": {"type": "string", "description": "e.g. A1:H20, max 400 cells"},
        }, "required": ["file_id", "sheet", "range"]},
    },
    {
        "name": "excel_lookup",
        "description": "Get one rate: give the row label (e.g. 'Up to 5 kg') and column header (e.g. 'Kanto'). "
                       "Returns the value and its cell for every matching table.",
        "parameters": {"type": "object", "properties": {
            "file_id": FILE_ID,
            "row": {"type": "string", "description": "row label, e.g. 'Up to 5 kg'"},
            "column": {"type": "string", "description": "column header, e.g. 'Kanto' or '4 ton Truck'"},
            "sheet": {"type": "string", "description": "sheet name (optional; all sheets if empty)"},
        }, "required": ["file_id", "row", "column"]},
    },
    {
        "name": "excel_search",
        "description": "Find cells containing a word, e.g. 'valid', 'fuel', 'minimum'.",
        "parameters": {"type": "object", "properties": {
            "file_id": FILE_ID,
            "text": {"type": "string", "description": "word or phrase"},
        }, "required": ["file_id", "text"]},
    },
]


class RunRequest(BaseModel):
    name: str
    arguments: dict = {}


class ToolError(Exception):
    pass


@app.get("/health")
def health():
    return {"ok": os.path.isdir(FILES_DIR)}


@app.get("/describe")
def describe():
    return {"service": "excel-tool", "tools": TOOLS}


@app.post("/run")
def run(req: RunRequest):
    a = req.arguments
    try:
        if req.name == "excel_profile":
            return {"ok": True, "result": profile(a.get("file_id", ""))}
        if req.name == "excel_read_range":
            return {"ok": True, "result": read_range(a.get("file_id", ""), a.get("sheet", ""),
                                                     a.get("range", ""))}
        if req.name == "excel_lookup":
            return {"ok": True, "result": lookup(a.get("file_id", ""), a.get("row", ""),
                                                 a.get("column", ""), a.get("sheet") or "")}
        if req.name == "excel_search":
            return {"ok": True, "result": search(a.get("file_id", ""), a.get("text", ""))}
        return {"ok": False, "error": f"excel-tool has no tool named '{req.name}'."}
    except ToolError as e:
        return {"ok": False, "error": str(e)}
    except Exception as e:  # noqa: BLE001 - a bad file must never crash the service
        return {"ok": False, "error": f"Could not read the file: {type(e).__name__}: {e}"}


# ---------------------------------------------------------------- file lookup

def _resolve(file_id: str) -> tuple[str, dict]:
    file_id = (file_id or "").strip()
    if not re.fullmatch(r"[0-9a-f]{32}", file_id):
        raise ToolError(f"'{file_id}' is not a valid file_id.")
    folder = os.path.join(FILES_DIR, file_id)
    try:
        with open(os.path.join(folder, "meta.json"), encoding="utf-8") as f:
            meta = json.load(f)
    except FileNotFoundError:
        raise ToolError(f"No attached file with file_id {file_id}.") from None
    return os.path.join(folder, meta["stored_name"]), meta


def _sheet(sheets: list[Sheet], name: str) -> Sheet:
    for s in sheets:
        if s.name == name:
            return s
    for s in sheets:
        if s.name.strip().lower() == name.strip().lower():
            return s
    raise ToolError(f"No sheet named '{name}'. Sheets: {', '.join(s.name for s in sheets)}")


# ---------------------------------------------------------------- profile

def _filled(row: list[object]) -> int:
    return sum(1 for v in row if v is not None and str(v).strip() != "")


def _guess_header(rows: list[list[object]]) -> int | None:
    """Index of the first 'wide' row in the top 30: titles above a table are narrow."""
    top = rows[:30]
    widest = max((_filled(r) for r in top), default=0)
    if widest < 2:
        return None
    for i, r in enumerate(top):
        if _filled(r) >= max(2, 0.6 * widest):
            return i
    return None


def _kind(v: object) -> str | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        return "text"
    if isinstance(v, (int, float)):
        return "number"
    if hasattr(v, "year"):
        return "date"
    return "text"


def _column_kinds(rows: list[list[object]], n_cols: int) -> str:
    """'A text, B-R number' style summary of what each column holds."""
    kinds = []
    for c in range(n_cols):
        counts: dict[str, int] = {}
        for r in rows:
            k = _kind(r[c]) if c < len(r) else None
            if k:
                counts[k] = counts.get(k, 0) + 1
        if not counts:
            kinds.append("empty")
        else:
            top, n = max(counts.items(), key=lambda kv: kv[1])
            kinds.append(top if n >= 0.8 * sum(counts.values()) else "mixed")
    parts, start = [], 0
    for c in range(1, n_cols + 1):
        if c == n_cols or kinds[c] != kinds[start]:
            a, b = get_column_letter(start + 1), get_column_letter(c)
            parts.append(f"{a if a == b else a + '-' + b} {kinds[start]}")
            start = c
    return ", ".join(parts)


def _row_text(row: list[object], n: int = MAX_COLS, width: int = 30) -> str:
    cells = [fmt(v, width) for v in row[:n]]
    more = f" | …(+{len(row) - n} cols)" if len(row) > n else ""
    return " | ".join(cells) + more


PROFILE_MAX_CHARS = 7500    # keeps the profile inside a small model's memory
MAX_LABELS = 30
MAX_TEXT_LINES = 8


def _text_line(row: list[object], width: int = 120) -> str:
    """A title/note row: only its non-empty cells."""
    return " | ".join(fmt(v, width) for v in row if v is not None and str(v).strip())


def _header_line(row: list[object]) -> str:
    """Header cells tagged with their column letter, so read_range can target them."""
    return " | ".join(f"{get_column_letter(i + 1)}:{fmt(v)}" for i, v in enumerate(row[:MAX_COLS])
                      if v is not None and str(v).strip())


def _first_value(row: list[object]) -> str:
    for v in row:
        if v is not None and str(v).strip():
            return fmt(v, 30)
    return ""


def _blocks(rows: list[list[object]]) -> list[tuple[int, int]]:
    """Runs of non-empty rows, as (first, last) indexes."""
    out, start = [], None
    for i, r in enumerate(rows):
        if _filled(r):
            if start is None:
                start = i
        elif start is not None:
            out.append((start, i - 1))
            start = None
    if start is not None:
        out.append((start, len(rows) - 1))
    return out


def _profile_sheet(s: Sheet) -> list[str]:
    rows = s.rows
    widest = max((_filled(r) for r in rows), default=0)
    wide = max(3, 0.5 * widest)
    lines: list[str] = []
    for a, b in _blocks(rows):
        idx = list(range(a, b + 1))
        wide_rows = [i for i in idx if _filled(rows[i]) >= wide]
        if not wide_rows:                       # titles or notes
            for i in idx[:MAX_TEXT_LINES]:
                lines.append(f"  r{i + 1}: {_text_line(rows[i])}")
            if len(idx) > MAX_TEXT_LINES:
                lines.append(f"  …(+{len(idx) - MAX_TEXT_LINES} more text rows)")
            continue
        h = wide_rows[0]
        if len(wide_rows) == 1 and not any(_filled(rows[i]) >= 2 for i in range(h + 1, b + 1)):
            for i in idx[:MAX_TEXT_LINES]:            # a lone wide row is a label line, not a table
                lines.append(f"  r{i + 1}: {_text_line(rows[i])}")
            continue
        for i in idx:
            if i >= h:
                break
            lines.append(f"  r{i + 1}: {_text_line(rows[i])}")
        last = wide_rows[-1]
        data = [i for i in range(h + 1, last + 1) if _filled(rows[i]) >= 2]
        lines.append(f"  TABLE r{h + 1}-r{last + 1}. Header: {_header_line(rows[h])}")
        if data:
            labels = [f"r{i + 1} {_first_value(rows[i])}" for i in data]
            more = f", …(+{len(labels) - MAX_LABELS})" if len(labels) > MAX_LABELS else ""
            lines.append(f"  Row labels ({len(labels)} rows): {', '.join(labels[:MAX_LABELS])}{more}")
            shown = data[:2] + ([data[-1]] if len(data) > 2 else [])
            for i in shown:
                lines.append(f"  r{i + 1}: {_row_text(rows[i])}")
        notes = [i for i in range(last + 1, b + 1)]
        for i in notes[:MAX_TEXT_LINES]:
            lines.append(f"  r{i + 1}: {_text_line(rows[i])}")
        if len(notes) > MAX_TEXT_LINES:
            lines.append(f"  …(+{len(notes) - MAX_TEXT_LINES} more text rows)")
    return lines


def profile(file_id: str) -> str:
    path, meta = _resolve(file_id)
    sheets = load(path, PROFILE_MAX_ROWS)
    out = [f"File: {meta['original_name']} ({meta.get('size', 0) / 1024:.0f} KB), "
           f"{len(sheets)} sheet(s): {', '.join(repr(s.name) for s in sheets)}",
           "Rows are numbered rN; header cells show their column letter. "
           "Use excel_read_range for exact values."]
    seen: dict[tuple, tuple[str, set[str]]] = {}   # table layout -> (sheet name, its lines)
    for s in sheets[:MAX_SHEETS]:
        out.append("")
        merged = ("" if s.merged is None else f", {len(s.merged)} merged range(s)")
        out.append(f"Sheet '{s.name}': {s.n_rows} rows x {s.n_cols} cols{merged}")
        if not any(_filled(r) for r in s.rows):
            out.append("  (empty)")
            continue
        lines = _profile_sheet(s)
        layout = tuple(l.split(":", 1)[1] for l in lines
                       if l.startswith("  TABLE") or l.startswith("  Row labels"))
        if layout and layout in seen:
            first, first_lines = seen[layout]
            diff = [l for l in lines if l.split(": ", 1)[-1] not in first_lines
                    and not l.startswith("  TABLE") and not l.startswith("  Row labels")]
            out.append(f"  Same table layout as sheet '{first}' (same headers and row labels; "
                       f"values differ - read them with excel_read_range on this sheet). Lines that differ:")
            out.extend(diff[:12])
            continue
        if layout:
            seen[layout] = (s.name, {l.split(": ", 1)[-1] for l in lines})
        out.extend(lines)
        if s.n_rows > PROFILE_MAX_ROWS:
            out.append(f"  (profile read the first {PROFILE_MAX_ROWS} of {s.n_rows} rows)")
    if len(sheets) > MAX_SHEETS:
        out.append(f"\n(+{len(sheets) - MAX_SHEETS} more sheets not profiled)")
    text = "\n".join(out)
    if len(text) > PROFILE_MAX_CHARS:
        text = text[:PROFILE_MAX_CHARS] + "\n…(profile cut to fit; use excel_read_range or excel_search for the rest)"
    return text


# ---------------------------------------------------------------- lookup

def _norm(v: object) -> str:
    return " ".join(str(v).lower().replace("\n", " ").split()) if v is not None else ""


def _tables(s: Sheet) -> list[dict]:
    """Every table in a sheet: header row, data rows and the title line just above it."""
    rows = s.rows
    widest = max((_filled(r) for r in rows), default=0)
    wide = max(3, 0.5 * widest)
    out, last_text = [], ""
    for a, b in _blocks(rows):
        wide_rows = [i for i in range(a, b + 1) if _filled(rows[i]) >= wide]
        title = last_text
        for i in range(a, (wide_rows[0] if wide_rows else b + 1)):
            last_text = _text_line(rows[i])
            title = last_text
        if len(wide_rows) < 2:
            if wide_rows:
                last_text = _text_line(rows[wide_rows[0]])
            continue
        h, end = wide_rows[0], wide_rows[-1]
        out.append({"title": title, "header": h, "data": [i for i in range(h + 1, end + 1)
                                                            if _filled(rows[i]) >= 2]})
    return out


def _label_columns(s: Sheet, t: dict) -> list[int]:
    """Columns that hold row labels (mostly text below the header), e.g. A, or A and E side by side."""
    cols = []
    width = max((len(s.rows[i]) for i in t["data"]), default=0)
    for c in range(width):
        vals = [s.rows[i][c] for i in t["data"] if c < len(s.rows[i]) and s.rows[i][c] not in (None, "")]
        if vals and sum(isinstance(v, str) for v in vals) >= 0.6 * len(vals):
            cols.append(c)
    return cols or [0]


def _match(wanted: str, have: str) -> int:
    """2 = same text, 1 = one contains the other, 0 = no match."""
    if not have:
        return 0
    if wanted == have:
        return 2
    return 1 if wanted in have or have in wanted else 0


def lookup(file_id: str, row: str, column: str, sheet: str = "") -> str:
    path, _ = _resolve(file_id)
    want_r, want_c = _norm(row), _norm(column)
    if not want_r or not want_c:
        raise ToolError("Give both a row label and a column header.")
    sheets = load(path, PROFILE_MAX_ROWS)
    if sheet:
        sheets = [_sheet(sheets, sheet)]
    hits = []
    for s in sheets:
        for t in _tables(s):
            header = s.rows[t["header"]]
            cols = [(c, _match(want_c, _norm(v))) for c, v in enumerate(header)]
            cols = [x for x in cols if x[1]]
            if not cols:
                continue
            best_c = max(m for _, m in cols)
            label_cols = _label_columns(s, t)
            for c, m in cols:
                if m != best_c or c in label_cols:
                    continue
                lc = max([x for x in label_cols if x < c], default=label_cols[0])
                rows_found = [(i, _match(want_r, _norm(s.rows[i][lc] if lc < len(s.rows[i]) else None)))
                              for i in t["data"]]
                rows_found = [x for x in rows_found if x[1]]
                if not rows_found:
                    continue
                best_r = max(m2 for _, m2 in rows_found)
                for i, m2 in rows_found:
                    if m2 != best_r:
                        continue
                    v = s.rows[i][c] if c < len(s.rows[i]) else None
                    cell = f"{get_column_letter(c + 1)}{i + 1}"
                    hits.append(f"'{s.name}'!{cell} = {fmt(v, 60) or '(empty)'}   "
                                f"[row '{fmt(s.rows[i][lc], 40)}', column '{fmt(header[c], 40)}'; "
                                f"table r{t['header'] + 1}: {fmt(t['title'], 90) or 'no title'}]")
    if not hits:
        return (f"No cell found for row '{row}' and column '{column}'. Check the exact labels in "
                f"excel_profile (row labels start with rN, headers show column letters).")
    note = ("\nSeveral tables match: choose by the table title (e.g. normal vs dangerous goods, "
            "or which departure point)." if len(hits) > 1 else "")
    return "\n".join(hits[:12]) + note


# ---------------------------------------------------------------- read / search

def read_range(file_id: str, sheet: str, cell_range: str) -> str:
    path, _ = _resolve(file_id)
    try:
        c1, r1, c2, r2 = range_boundaries(cell_range.strip().upper())
    except (ValueError, TypeError):
        raise ToolError(f"'{cell_range}' is not a valid range. Use something like A1:H20.") from None
    if None in (c1, r1, c2, r2):
        raise ToolError("Give a full range with both corners, like A1:H20.")
    cells = (c2 - c1 + 1) * (r2 - r1 + 1)
    if cells > MAX_RANGE_CELLS:
        raise ToolError(f"Range has {cells} cells; the limit is {MAX_RANGE_CELLS}. Ask for a smaller range.")
    s = _sheet(load(path, r2), sheet)
    lines = [f"Sheet '{s.name}' {cell_range.upper()}:",
             "     " + " | ".join(get_column_letter(c) for c in range(c1, c2 + 1))]
    for r in range(r1, r2 + 1):
        row = s.rows[r - 1] if r - 1 < len(s.rows) else []
        vals = [fmt(row[c - 1], 40) if c - 1 < len(row) else "" for c in range(c1, c2 + 1)]
        lines.append(f"r{r}: " + " | ".join(vals))
    return "\n".join(lines)


def search(file_id: str, text: str) -> str:
    path, _ = _resolve(file_id)
    needle = (text or "").strip().lower()
    if len(needle) < 2:
        raise ToolError("Search text must be at least 2 characters.")
    hits, total = [], 0
    for s in load(path, 5000):
        for ri, row in enumerate(s.rows):
            for ci, v in enumerate(row):
                if v is not None and needle in str(v).lower():
                    total += 1
                    if len(hits) < MAX_SEARCH_HITS:
                        hits.append(f"'{s.name}'!{get_column_letter(ci + 1)}{ri + 1}: {fmt(v, 80)}")
    if not hits:
        return f"No cells contain '{text}' (searched the first 5000 rows of each sheet)."
    more = f"\n(+{total - len(hits)} more matches)" if total > len(hits) else ""
    return f"{total} match(es) for '{text}':\n" + "\n".join(hits) + more
