"""Opens spreadsheets safely and gives every format the same simple shape.

Safety: formulas are never calculated (cached values only), external workbook
links are not loaded, macros are not loaded, and nothing is ever fetched.
"""
import csv
import datetime as dt
import os
from dataclasses import dataclass, field

MAX_FULL_LOAD_BYTES = 15 * 1024 * 1024   # bigger .xlsx files load read-only (no merged-cell info)


@dataclass
class Sheet:
    name: str
    rows: list[list[object]]          # cell values, row by row (trailing empties trimmed)
    n_rows: int
    n_cols: int
    merged: list[str] | None = field(default=None)   # None = unknown for this file


def fmt(v: object, width: int = 30) -> str:
    """Short, readable text for one cell."""
    if v is None:
        return ""
    if isinstance(v, bool):
        s = "TRUE" if v else "FALSE"
    elif isinstance(v, float):
        s = f"{v:.4f}".rstrip("0").rstrip(".") if v != int(v) else str(int(v))
    elif isinstance(v, dt.datetime):
        s = v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ", timespec="minutes")
    elif isinstance(v, dt.date):
        s = v.isoformat()
    else:
        s = " ".join(str(v).split())
    return s if len(s) <= width else s[: width - 1] + "…"


def _trim(row: list[object]) -> list[object]:
    while row and (row[-1] is None or (isinstance(row[-1], str) and not row[-1].strip())):
        row.pop()
    return row


def load(path: str, max_rows: int) -> list[Sheet]:
    """Every sheet, with up to `max_rows` rows of values each."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _load_openpyxl(path, max_rows)
    if ext == ".xls":
        return _load_xlrd(path, max_rows)
    if ext == ".csv":
        return _load_csv(path, max_rows)
    raise ValueError(f"Unsupported file type '{ext}'. Supported: .xlsx .xlsm .xls .csv")


def _load_openpyxl(path: str, max_rows: int) -> list[Sheet]:
    import openpyxl
    read_only = os.path.getsize(path) > MAX_FULL_LOAD_BYTES
    wb = openpyxl.load_workbook(path, read_only=read_only, data_only=True,
                                keep_links=False, keep_vba=False)
    try:
        sheets = []
        for ws in wb.worksheets:
            rows = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i >= max_rows:
                    break
                rows.append(_trim(list(r)))
            merged = None if read_only else [str(m) for m in ws.merged_cells.ranges]
            sheets.append(Sheet(ws.title, rows, ws.max_row or 0, ws.max_column or 0, merged))
        return sheets
    finally:
        wb.close()


def _load_xlrd(path: str, max_rows: int) -> list[Sheet]:
    import xlrd
    try:
        book = xlrd.open_workbook(path, formatting_info=True, on_demand=True)
        with_merged = True
    except (NotImplementedError, xlrd.XLRDError):
        book = xlrd.open_workbook(path, on_demand=True)
        with_merged = False
    try:
        sheets = []
        for sh in book.sheets():
            rows = []
            for r in range(min(sh.nrows, max_rows)):
                vals = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    v = cell.value
                    if cell.ctype == xlrd.XL_CELL_EMPTY or cell.ctype == xlrd.XL_CELL_BLANK:
                        v = None
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                        except Exception:  # noqa: BLE001
                            pass
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        v = "#ERR"
                    vals.append(v)
                rows.append(_trim(vals))
            merged = None
            if with_merged:
                from openpyxl.utils import get_column_letter as L
                merged = [f"{L(c1 + 1)}{r1 + 1}:{L(c2)}{r2}" for r1, r2, c1, c2 in sh.merged_cells]
            sheets.append(Sheet(sh.name, rows, sh.nrows, sh.ncols, merged))
        return sheets
    finally:
        book.release_resources()


def _load_csv(path: str, max_rows: int) -> list[Sheet]:
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
        sample = f.read(64 * 1024)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows, total, width = [], 0, 0
        for r in csv.reader(f, dialect):
            total += 1
            width = max(width, len(r))
            if len(rows) < max_rows:
                rows.append(_trim([_num(x) for x in r]))
    return [Sheet(os.path.basename(path), rows, total, width, [])]


def _num(s: str) -> object:
    s = s.strip()
    if not s:
        return None
    try:
        return float(s.replace(",", "")) if any(ch.isdigit() for ch in s) else s
    except ValueError:
        return s
